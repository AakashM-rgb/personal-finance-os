import csv
import io
from datetime import date

from httpx import AsyncClient
from openpyxl import load_workbook
from pypdf import PdfReader


def _iso(d: date, time_str: str = "09:00:00") -> str:
    return f"{d.isoformat()}T{time_str}Z"


async def _create_account(client: AsyncClient, headers: dict, **payload) -> dict:
    payload.setdefault("name", "Test Account")
    payload.setdefault("type", "cash")
    payload.setdefault("balance_minor", 10000000)
    response = await client.post("/api/v1/accounts", headers=headers, json=payload)
    assert response.status_code == 201, response.text
    return response.json()["data"]


async def _get_category_id(client: AsyncClient, headers: dict, name: str) -> str:
    response = await client.get("/api/v1/categories", headers=headers)
    for category in response.json()["data"]:
        if category["name"] == name:
            return category["id"]
    raise AssertionError(f"category {name!r} not found in list")


async def _create_transaction(client: AsyncClient, headers: dict, **payload) -> dict:
    response = await client.post("/api/v1/transactions", headers=headers, json=payload)
    assert response.status_code == 201, response.text
    return response.json()["data"]


async def _seed_category_data(client: AsyncClient, headers: dict) -> None:
    account = await _create_account(client, headers)
    food_id = await _get_category_id(client, headers, "Food")
    await _create_transaction(
        client,
        headers,
        account_id=account["id"],
        type="expense",
        amount_minor=324000,
        category_id=food_id,
        occurred_at=_iso(date(2026, 6, 15)),
        description="Groceries",
    )


# --- CSV -----------------------------------------------------------------------------


async def test_category_report_csv_export_contains_real_data(
    client: AsyncClient, auth_headers: dict
) -> None:
    await _seed_category_data(client, auth_headers)

    response = await client.get(
        "/api/v1/reports/category/export",
        headers=auth_headers,
        params={
            "format": "csv",
            "range": "custom",
            "custom_from": "2026-06-01",
            "custom_to": "2026-06-30",
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith("text/csv")
    assert "attachment" in response.headers["content-disposition"]

    content = response.content.decode("utf-8-sig")
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    assert any("Food" in row for row in rows)
    assert any("3240.00 INR" in row for row in rows)


async def test_category_report_csv_export_empty_data_has_no_fake_rows(
    client: AsyncClient, auth_headers: dict
) -> None:
    response = await client.get(
        "/api/v1/reports/category/export",
        headers=auth_headers,
        params={"format": "csv", "range": "current_month"},
    )
    assert response.status_code == 200
    content = response.content.decode("utf-8-sig")
    assert "No data available for this range." in content


# --- Excel ---------------------------------------------------------------------------


async def test_budget_report_excel_export_is_valid_and_contains_real_data(
    client: AsyncClient, auth_headers: dict
) -> None:
    food_id = await _get_category_id(client, auth_headers, "Food")
    await client.post(
        "/api/v1/budgets",
        headers=auth_headers,
        json={"category_id": food_id, "amount_minor": 500000},
    )

    response = await client.get(
        "/api/v1/reports/budget/export", headers=auth_headers, params={"format": "xlsx"}
    )
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument"
    )

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    all_rows = [r for r in sheet.iter_rows(values_only=True) if any(c is not None for c in r)]
    assert any("Food" in str(cell) for row in all_rows for cell in row)
    assert any(cell == 5000.0 for row in all_rows for cell in row)  # 500000 minor -> 5000.00


async def test_net_worth_report_excel_export_empty_accounts(
    client: AsyncClient, auth_headers: dict
) -> None:
    response = await client.get(
        "/api/v1/reports/net-worth/export", headers=auth_headers, params={"format": "xlsx"}
    )
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.content))
    assert workbook.active is not None


# --- PDF -----------------------------------------------------------------------------


async def test_income_report_pdf_export_contains_title_and_summary(
    client: AsyncClient, auth_headers: dict
) -> None:
    account = await _create_account(client, auth_headers)
    await _create_transaction(
        client,
        auth_headers,
        account_id=account["id"],
        type="income",
        amount_minor=500000,
        occurred_at=_iso(date(2026, 6, 1)),
        description="Salary",
    )

    response = await client.get(
        "/api/v1/reports/income/export",
        headers=auth_headers,
        params={
            "format": "pdf",
            "range": "custom",
            "custom_from": "2026-06-01",
            "custom_to": "2026-06-30",
        },
    )
    assert response.status_code == 200
    assert response.headers["content-type"] == "application/pdf"
    assert response.content[:4] == b"%PDF"

    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "Income Report" in text
    assert "5000.00 INR" in text


async def test_monthly_report_pdf_export_empty_month(
    client: AsyncClient, auth_headers: dict
) -> None:
    response = await client.get(
        "/api/v1/reports/monthly/export",
        headers=auth_headers,
        params={"format": "pdf", "year": 2026, "month": 6},
    )
    assert response.status_code == 200
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(response.content)).pages)
    assert "Monthly Report" in text
    assert "No data available for this range." in text


# --- authorization isolation ------------------------------------------------------------


async def test_export_never_leaks_another_users_data(
    client: AsyncClient, auth_headers: dict, other_auth_headers: dict
) -> None:
    await _seed_category_data(client, other_auth_headers)

    response = await client.get(
        "/api/v1/reports/category/export",
        headers=auth_headers,  # a DIFFERENT user requests the export
        params={
            "format": "csv",
            "range": "custom",
            "custom_from": "2026-06-01",
            "custom_to": "2026-06-30",
        },
    )
    assert response.status_code == 200
    content = response.content.decode("utf-8-sig")
    assert "Food" not in content
    assert "3240.00" not in content
    assert "No data available for this range." in content


async def test_export_requires_authentication(client: AsyncClient) -> None:
    response = await client.get(
        "/api/v1/reports/category/export", params={"format": "csv", "range": "current_month"}
    )
    assert response.status_code == 401
