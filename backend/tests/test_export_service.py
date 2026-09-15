import csv
import io
from datetime import date

from openpyxl import load_workbook
from pypdf import PdfReader

from app.services.export_service import (
    ExportColumn,
    ExportTable,
    render_csv,
    render_excel,
    render_pdf,
)

_SAMPLE_TABLE = ExportTable(
    title="Category Report",
    subtitle="2026-09-01 to 2026-09-30",
    currency="INR",
    summary=[("Total expenses", "1,950.00 INR")],
    columns=[
        ExportColumn("Category", "text"),
        ExportColumn("Amount", "money"),
        ExportColumn("Percent", "percent"),
        ExportColumn("Count", "int"),
    ],
    rows=[
        ["Food", 175000, 89.7, 3],
        ["Uncategorized", 20000, 10.3, 1],
    ],
)

_EMPTY_TABLE = ExportTable(
    title="Category Report",
    subtitle="2026-09-01 to 2026-09-30",
    currency="INR",
    summary=[("Total expenses", "0.00 INR")],
    columns=[ExportColumn("Category", "text"), ExportColumn("Amount", "money")],
    rows=[],
)

_ESCAPING_TABLE = ExportTable(
    title="Category Report",
    subtitle="range",
    currency="INR",
    summary=[],
    columns=[ExportColumn("Category", "text"), ExportColumn("Amount", "money")],
    rows=[
        ['Coffee, Tea "& Snacks"', 5000],
        ["Rent\nUtilities", 300000],
    ],
)


# --- CSV -----------------------------------------------------------------------------


def test_csv_contains_title_and_summary() -> None:
    content = render_csv(_SAMPLE_TABLE).decode("utf-8-sig")
    assert "Category Report" in content
    assert "Total expenses" in content
    assert "1,950.00 INR" in content


def test_csv_has_correct_headers_and_rows() -> None:
    content = render_csv(_SAMPLE_TABLE).decode("utf-8-sig")
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    header_row = next(r for r in rows if r == ["Category", "Amount", "Percent", "Count"])
    header_index = rows.index(header_row)
    data_rows = rows[header_index + 1 :]
    assert data_rows[0] == ["Food", "1750.00 INR", "89.7%", "3"]
    assert data_rows[1] == ["Uncategorized", "200.00 INR", "10.3%", "1"]


def test_csv_escapes_commas_and_quotes_correctly() -> None:
    content = render_csv(_ESCAPING_TABLE).decode("utf-8-sig")
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    # A correctly-escaped CSV round-trips through the csv module back to
    # the exact original string, commas/quotes/newlines and all.
    assert ['Coffee, Tea "& Snacks"', "50.00 INR"] in rows
    assert ["Rent\nUtilities", "3000.00 INR"] in rows


def test_csv_empty_data_shows_a_clear_message_not_fake_rows() -> None:
    content = render_csv(_EMPTY_TABLE).decode("utf-8-sig")
    assert "No data available for this range." in content
    reader = csv.reader(io.StringIO(content))
    rows = list(reader)
    # No row looks like fabricated data - only text/header/summary rows.
    for row in rows:
        assert row == [] or all(not cell.replace(".", "").isdigit() for cell in row if cell)


# --- Excel ---------------------------------------------------------------------------


def test_excel_is_a_valid_loadable_workbook() -> None:
    content = render_excel(_SAMPLE_TABLE)
    workbook = load_workbook(io.BytesIO(content))
    assert workbook.sheetnames == ["Category Report"]


def test_excel_has_correct_headers_and_numeric_types() -> None:
    workbook = load_workbook(io.BytesIO(render_excel(_SAMPLE_TABLE)))
    sheet = workbook.active
    all_rows = list(sheet.iter_rows(values_only=True))
    header_row_index = next(i for i, r in enumerate(all_rows) if r[0] == "Category")
    header = all_rows[header_row_index]
    assert header == ("Category", "Amount", "Percent", "Count")

    food_row = all_rows[header_row_index + 1]
    assert food_row[0] == "Food"
    # openpyxl round-trips a whole-number float (1750.0) back as a plain
    # int - either way it is a real number, never a formatted string.
    assert isinstance(food_row[1], int | float)
    assert not isinstance(food_row[1], str)
    assert float(food_row[1]) == 1750.0  # 175000 minor units -> 1750.00 major
    assert isinstance(food_row[2], int | float)
    assert float(food_row[2]) == 89.7
    assert isinstance(food_row[3], int)
    assert food_row[3] == 3


def test_excel_money_column_has_a_number_format() -> None:
    workbook = load_workbook(io.BytesIO(render_excel(_SAMPLE_TABLE)))
    sheet = workbook.active
    all_rows = list(sheet.iter_rows())
    header_row_index = next(i for i, r in enumerate(all_rows) if r[0].value == "Category")
    amount_cell = all_rows[header_row_index + 1][1]
    assert amount_cell.number_format == "#,##0.00"


def test_excel_date_column_uses_a_real_date_type() -> None:
    table = ExportTable(
        title="Yearly Report",
        subtitle="2026",
        currency="INR",
        summary=[],
        columns=[ExportColumn("Month", "date"), ExportColumn("Income", "money")],
        rows=[[date(2026, 1, 1), 500000]],
    )
    workbook = load_workbook(io.BytesIO(render_excel(table)))
    sheet = workbook.active
    all_rows = list(sheet.iter_rows(values_only=True))
    header_row_index = next(i for i, r in enumerate(all_rows) if r[0] == "Month")
    month_value = all_rows[header_row_index + 1][0]
    assert month_value.year == 2026
    assert month_value.month == 1
    assert month_value.day == 1


def test_excel_sheet_name_is_truncated_to_31_characters() -> None:
    long_title = "A" * 50
    table = ExportTable(
        title=long_title,
        subtitle="sub",
        currency="INR",
        summary=[],
        columns=[ExportColumn("X", "text")],
        rows=[],
    )
    workbook = load_workbook(io.BytesIO(render_excel(table)))
    assert len(workbook.sheetnames[0]) <= 31


def test_excel_empty_data_has_no_fake_rows() -> None:
    workbook = load_workbook(io.BytesIO(render_excel(_EMPTY_TABLE)))
    sheet = workbook.active
    all_rows = [r for r in sheet.iter_rows(values_only=True) if any(r)]
    assert any("No data available" in str(cell) for row in all_rows for cell in row if cell)
    # No numeric-looking data row beyond the header exists.
    header_row_index = next(i for i, r in enumerate(all_rows) if r[0] == "Category")
    assert "No data" in str(all_rows[header_row_index + 1][0])


# --- PDF -----------------------------------------------------------------------------


def test_pdf_is_valid_and_readable() -> None:
    content = render_pdf(_SAMPLE_TABLE)
    assert content[:4] == b"%PDF"
    reader = PdfReader(io.BytesIO(content))
    assert len(reader.pages) >= 1


def test_pdf_contains_title_range_and_summary() -> None:
    content = render_pdf(_SAMPLE_TABLE)
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)
    assert "Category Report" in text
    assert "2026-09-01 to 2026-09-30" in text
    assert "Total expenses" in text
    assert "1,950.00 INR" in text


def test_pdf_contains_table_rows() -> None:
    content = render_pdf(_SAMPLE_TABLE)
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)
    assert "Food" in text
    assert "1750.00 INR" in text
    assert "Uncategorized" in text


def test_pdf_empty_data_shows_a_clear_message() -> None:
    content = render_pdf(_EMPTY_TABLE)
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(content)).pages)
    assert "No data available for this range." in text


def test_pdf_paginates_a_large_table_without_clipping_content() -> None:
    many_rows = [[f"Category {i}", 1000 * i] for i in range(200)]
    table = ExportTable(
        title="Category Report",
        subtitle="a big range",
        currency="INR",
        summary=[],
        columns=[ExportColumn("Category", "text"), ExportColumn("Amount", "money")],
        rows=many_rows,
    )
    content = render_pdf(table)
    reader = PdfReader(io.BytesIO(content))
    assert len(reader.pages) > 1  # a 200-row table must spill onto more than one page
    full_text = "\n".join(page.extract_text() for page in reader.pages)
    assert "Category 0" in full_text
    assert "Category 199" in full_text  # the very last row must not be clipped off
