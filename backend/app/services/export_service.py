"""Export rendering: turns an ExportTable (a small, report-agnostic
tabular representation) into real CSV, Excel (.xlsx), or PDF bytes.
Every ExportTable is built directly from a real report response (see
app.services.report_export_service) - this module only renders, it never
invents rows, and an empty report renders as a valid file with zero data
rows plus a clear "no data" note, never fabricated content.
"""

import csv
import io
from dataclasses import dataclass
from datetime import date, datetime
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.services.export_formatting import format_money, minor_to_major

ExportFormat = Literal["csv", "xlsx", "pdf"]
ColumnKind = Literal["text", "money", "percent", "int", "date"]

CONTENT_TYPES: dict[ExportFormat, str] = {
    "csv": "text/csv",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}


@dataclass(frozen=True)
class ExportColumn:
    header: str
    kind: ColumnKind


@dataclass(frozen=True)
class ExportTable:
    """A report-agnostic tabular export payload. `rows` holds each
    column's RAW value (int minor units for "money", float 0-100 for
    "percent", a real `date` for "date", else a plain value) - each
    renderer formats them appropriately for its own file format, so a
    money column is a real number in Excel, not a pre-formatted string."""

    title: str
    subtitle: str
    currency: str
    summary: list[tuple[str, str]]
    columns: list[ExportColumn]
    rows: list[list[object]]


def _cell_display(value: object, kind: ColumnKind, currency: str) -> str:
    if kind == "money":
        assert isinstance(value, int)
        return format_money(value, currency)
    if kind == "percent":
        assert isinstance(value, int | float)
        return f"{float(value):.1f}%"
    if kind == "date":
        return value.isoformat() if isinstance(value, date) else str(value)
    return str(value)


def render_csv(table: ExportTable) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)

    writer.writerow([table.title])
    writer.writerow([table.subtitle])
    writer.writerow([])
    for label, value in table.summary:
        writer.writerow([label, value])
    writer.writerow([])

    writer.writerow([c.header for c in table.columns])
    if not table.rows:
        writer.writerow(["No data available for this range."])
    for row in table.rows:
        writer.writerow(
            [
                _cell_display(value, col.kind, table.currency)
                for value, col in zip(row, table.columns, strict=True)
            ]
        )

    # utf-8-sig (a leading BOM) so Excel on Windows recognizes the file as
    # UTF-8 instead of guessing a legacy codepage for any non-ASCII text
    # (category names, currency symbols) in the cells.
    return buffer.getvalue().encode("utf-8-sig")


def render_excel(table: ExportTable) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    # Excel sheet names are capped at 31 characters and can't contain
    # some punctuation - keep it simple and safe.
    sheet.title = table.title[:31].replace("/", "-").replace("\\", "-")

    bold = Font(bold=True)
    sheet.append([table.title])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([table.subtitle])
    sheet.append([])
    for label, value in table.summary:
        sheet.append([label, value])
    sheet.append([])

    header_row_index = sheet.max_row + 1
    sheet.append([c.header for c in table.columns])
    for cell in sheet[header_row_index]:
        cell.font = bold

    if not table.rows:
        sheet.append(["No data available for this range."])
    for row in table.rows:
        rendered_row: list[object] = []
        for cell_value, col in zip(row, table.columns, strict=True):
            if col.kind == "money":
                assert isinstance(cell_value, int)
                rendered_row.append(minor_to_major(cell_value, table.currency))
            elif col.kind == "percent":
                assert isinstance(cell_value, int | float)
                rendered_row.append(float(cell_value))
            elif col.kind == "date" and isinstance(cell_value, date):
                rendered_row.append(datetime(cell_value.year, cell_value.month, cell_value.day))
            else:
                rendered_row.append(cell_value)
        sheet.append(rendered_row)

        row_index = sheet.max_row
        for col_index, col in enumerate(table.columns, start=1):
            cell = sheet.cell(row=row_index, column=col_index)
            if col.kind == "money":
                decimals = 0 if table.currency == "JPY" else 2
                cell.number_format = "#,##0" if decimals == 0 else "#,##0.00"
            elif col.kind == "percent":
                cell.number_format = "0.0"
            elif col.kind == "date":
                cell.number_format = "yyyy-mm-dd"

    for col_index, col in enumerate(table.columns, start=1):
        sheet.column_dimensions[get_column_letter(col_index)].width = max(len(col.header) + 2, 14)

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def render_pdf(table: ExportTable) -> bytes:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=0.6 * inch,
        rightMargin=0.6 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
    )
    styles = getSampleStyleSheet()
    subtitle_style = ParagraphStyle(
        "Subtitle", parent=styles["Normal"], textColor=colors.grey, spaceAfter=6
    )

    elements = [
        Paragraph(table.title, styles["Title"]),
        Paragraph(table.subtitle, subtitle_style),
        Spacer(1, 0.1 * inch),
    ]

    if table.summary:
        summary_data = [[label, value] for label, value in table.summary]
        summary_table = Table(summary_data, colWidths=[2.2 * inch, 2.2 * inch])
        summary_table.setStyle(
            TableStyle(
                [
                    ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        elements.append(summary_table)
        elements.append(Spacer(1, 0.2 * inch))

    header = [c.header for c in table.columns]
    if table.rows:
        data_rows = [
            [
                _cell_display(value, col.kind, table.currency)
                for value, col in zip(row, table.columns, strict=True)
            ]
            for row in table.rows
        ]
        table_data = [header, *data_rows]
    else:
        # One row spanning every column, merged below, so the "no data"
        # message reads as a single sentence rather than a fragmented row.
        empty_row: list[str] = ["No data available for this range."] + [""] * (
            len(table.columns) - 1
        )
        table_data = [header, empty_row]

    data_table = Table(table_data, repeatRows=1)
    style_commands = [
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f0efec")),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#c3c2b7")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fcfcfb")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]
    if not table.rows:
        style_commands.append(("SPAN", (0, 1), (-1, 1)))
    data_table.setStyle(TableStyle(style_commands))
    elements.append(data_table)

    doc.build(elements)
    return buffer.getvalue()


def render(table: ExportTable, export_format: ExportFormat) -> bytes:
    if export_format == "csv":
        return render_csv(table)
    if export_format == "xlsx":
        return render_excel(table)
    return render_pdf(table)
